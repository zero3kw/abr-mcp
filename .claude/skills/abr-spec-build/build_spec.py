"""公式「データフォーマット（仕様確定版）」xlsx を変換して、各データセットの
項目定義を Markdown（specs/…md）として生成する。元の xlsx を正とする。

依存: openpyxl（`pip install openpyxl`）。MCPサーバの実行時には不要。

  python3 build_spec.py             # 公式URLからDLして変換
  python3 build_spec.py --xlsx F    # ローカルxlsxから変換
  python3 build_spec.py -o OUT.md   # 出力先を指定
"""

import argparse
import io
import urllib.request
from pathlib import Path

import openpyxl

URL = ('https://www.digital.go.jp/assets/contents/node/basic_page/field_ref_resources/'
       '3d996df3-f87c-4439-b474-c0fbbf3ddad8/31267348/'
       '20240115_policies_base_registry_format_01.xlsx')

# 既定の出力先（このスクリプトはリポジトリの .claude/skills/abr-spec-build/ にある想定）。
DEFAULT_OUT = (Path(__file__).resolve().parents[3]
               / 'specs' / '20240115_policies_base_registry_format_01.md')

# シート名 → ABR 物理テーブル名。xlsx には物理名が無いため対応表で補う
# （いずれも公式のファイル名／カラム接頭辞で使われている正規のテーブル名）。
# ここに無いシート（表紙・改訂履歴・説明）は出力しない。
IDS = {
    '01都道府県': 'mt_pref',
    '02市区町村': 'mt_city',
    '03町字': 'mt_town',
    '04地番': 'mt_parcel',
    '05住居表示-街区': 'mt_rsdtdsp_blk',
    '06住居表示-住居': 'mt_rsdtdsp_rsdt',
    '07町字フルセット': 'mt_town_fullset',
    '08都道府県位置参照': 'mt_pref_pos',
    '09市区町村位置参照': 'mt_city_pos',
    '10町字位置参照': 'mt_town_pos',
    '11地番位置参照': 'mt_parcel_pos',
    '12住居表示-街区位置参照': 'mt_rsdtdsp_blk_pos',
    '13住居表示-住居位置参照': 'mt_rsdtdsp_rsdt_pos',
}

# 出力する列（xlsx の列1〜8）。列9以降の「収録例」は出力しない。
HEADERS = ['項目No.', '項目名', '項目名(日本語表記)', 'データ型', '桁数', 'Not Null', 'Key', '説明']


def cell(v):
    """セル値を1行のテキストに均す（改行・全角空白を半角空白に、連続空白を1つに）。"""
    if v is None:
        return ''
    s = str(v).replace('\r', ' ').replace('\n', ' ').replace('　', ' ')
    return ' '.join(s.split()).replace('|', r'\|')  # 表を壊さないよう | をエスケープ


def header_row(ws):
    """『項目名』が入った見出し行の番号を返す。"""
    for r in range(1, min(ws.max_row, 12) + 1):
        if cell(ws.cell(r, 2).value) == '項目名':
            return r
    raise ValueError(f'見出し行が見つかりません: {ws.title}')


def sheet_rows(ws):
    """データ行（項目No. が数値の行）を列1〜8で取り出す。"""
    rows = []
    for r in range(header_row(ws) + 1, ws.max_row + 1):
        no = ws.cell(r, 1).value
        name = cell(ws.cell(r, 2).value)
        if not isinstance(no, (int, float)) or isinstance(no, bool) or not name:
            break  # 項目No. が切れたら（注記・変更点の行）終了
        rows.append([str(int(no))] + [cell(ws.cell(r, c).value) for c in range(2, 9)])
    return rows


def md_table(rows):
    lines = ['| ' + ' | '.join(HEADERS) + ' |',
             '| ' + ' | '.join('---' for _ in HEADERS) + ' |']
    lines += ['| ' + ' | '.join(cells) + ' |' for cells in rows]
    return '\n'.join(lines)


def build(wb):
    parts = ['---\n'
             f'source: {URL}\n'
             'note: abr-spec-build スキルが公式xlsxから自動生成。手動編集しないこと。\n'
             '---']
    for ws in wb.worksheets:
        if ws.title not in IDS:
            continue
        parts.append(f'# {ws.title}({IDS[ws.title]})\n\n' + md_table(sheet_rows(ws)))
    return '\n\n'.join(parts) + '\n'


def main():
    ap = argparse.ArgumentParser(description='ABR データフォーマット xlsx → Markdown')
    ap.add_argument('--xlsx', help='ローカル xlsx（無指定なら公式URLからダウンロード）')
    ap.add_argument('-o', '--out', default=str(DEFAULT_OUT), help='出力先 md')
    args = ap.parse_args()

    if args.xlsx:
        src = args.xlsx
    else:
        print(f'ダウンロード: {URL}')
        src = io.BytesIO(urllib.request.urlopen(URL).read())
    wb = openpyxl.load_workbook(src, data_only=True)
    Path(args.out).write_text(build(wb), encoding='utf-8')
    n = sum(1 for w in wb.worksheets if w.title in IDS)
    print(f'生成: {args.out}（{n} データセット）')


if __name__ == '__main__':
    main()
